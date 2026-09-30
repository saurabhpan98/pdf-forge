"""
Pre-conversion font audit for Word → PDF.

Extracts every font referenced by the DOCX and compares against the
fonts installed on the server. Reports which fonts will be substituted.

Output format on stdout:
    __RESULT__{"fonts":[...],"missing":[...],...}
"""
import sys
import json
import subprocess

MARKER = '__RESULT__'


def log(msg):
    sys.stderr.write(f"[fontaudit] {msg}\n")


def _emit(payload):
    sys.stdout.write(f"{MARKER}{json.dumps(payload)}\n")
    sys.stdout.flush()


def _silence_library_noise():
    try:
        import fitz
        try:
            fitz.TOOLS.mupdf_display_errors(False)
        except Exception:
            try:
                fitz.TOOLS.reset_mupdf_warnings()
            except Exception:
                pass
    except Exception:
        pass


def get_installed_fonts():
    try:
        result = subprocess.run(
            ['fc-list', '--format=%{family}\n'],
            capture_output=True, text=True, timeout=15,
        )
    except Exception as e:
        log(f"fc-list failed: {e}")
        return set()

    families = set()
    for line in result.stdout.split('\n'):
        for fam in line.split(','):
            fam = fam.strip().lower()
            if fam:
                families.add(fam)
    return families


# ---------------------------------------------------------------------------
# Substitution map.
#
# Key:    lowercase substring to match against the requested font name.
# Value: list of installed font names — if ANY is present, the font is
#        considered "available" (LibreOffice will substitute silently).
#        An empty list means "no reliable substitute exists — always warn."
# ---------------------------------------------------------------------------
SUBSTITUTIONS = {
    # ---- Microsoft Office core ----
    'calibri':          ['carlito'],
    'cambria':          ['caladea'],
    'aptos':            ['carlito'],        # new Word default (2023+)
    'aptos display':    ['carlito'],
    'aptos narrow':     ['carlito', 'liberation sans narrow'],
    'aptos mono':       ['liberation mono', 'dejavu sans mono'],
    'arial':            ['liberation sans', 'arimo'],
    'arial narrow':     ['liberation sans narrow'],
    'arial black':      ['liberation sans', 'archivo black'],
    'arial rounded':    ['liberation sans'],
    'helvetica':        ['liberation sans', 'arimo'],
    'helvetica neue':   ['liberation sans', 'arimo'],
    'times new roman':  ['liberation serif', 'tinos'],
    'times':            ['liberation serif', 'tinos'],
    'courier new':      ['liberation mono', 'cousine'],
    'courier':          ['liberation mono', 'cousine'],
    'verdana':          ['dejavu sans'],
    'tahoma':           ['dejavu sans'],
    'trebuchet ms':     ['dejavu sans'],
    'trebuchet':        ['dejavu sans'],
    'segoe ui':         ['carlito', 'open sans', 'liberation sans'],
    'segoe ui light':   ['carlito', 'open sans'],
    'segoe ui semibold':['carlito', 'open sans'],
    'segoe ui symbol':  ['noto sans symbols'],
    'segoe ui emoji':   ['noto color emoji'],
    'georgia':          ['caladea', 'gelasio', 'liberation serif'],
    'constantia':       ['caladea'],
    'candara':          ['carlito', 'open sans'],
    'corbel':           ['carlito', 'open sans'],
    'consolas':         ['hack', 'liberation mono', 'dejavu sans mono'],
    'comic sans ms':    [],   # no good match — warn every time
    'comic sans':       [],
    'franklin gothic':  ['liberation sans'],
    'franklin gothic book': ['liberation sans'],
    'franklin gothic medium': ['liberation sans'],
    'bookman old style':['liberation serif', 'urw bookman'],
    'book antiqua':     ['urw palladio', 'liberation serif'],
    'palatino':         ['urw palladio', 'liberation serif'],
    'palatino linotype':['urw palladio', 'liberation serif'],
    'century gothic':   ['urw gothic', 'jost'],
    'century schoolbook': ['urw bookman', 'c059'],
    'garamond':         ['eb garamond', 'urw garamond'],
    'minion':           ['eb garamond'],
    'wingdings':        [],   # no substitute — warn
    'wingdings 2':      [],
    'wingdings 3':      [],
    'webdings':         [],
    'symbol':           ['standard symbols ps', 'symbola'],
    'zapf dingbats':    ['dingbats'],

    # ---- Google / modern ---- 
    'roboto':           ['roboto'],
    'open sans':        ['open sans', 'carlito'],
    'lato':             ['lato', 'carlito'],
    'montserrat':       ['montserrat'],
    'poppins':          ['poppins', 'carlito'],
    'inter':            ['inter', 'carlito'],
    'nunito':           ['nunito', 'carlito'],
    'raleway':          ['raleway', 'carlito'],
    'work sans':        ['work sans', 'carlito'],
    'ubuntu':           ['ubuntu', 'carlito'],
    'rubik':            ['rubik', 'carlito'],
    'karla':            ['karla', 'carlito'],
    'mulish':           ['mulish', 'carlito'],
    'manrope':          ['manrope', 'carlito'],
    'dm sans':          ['dm sans', 'carlito'],
    'jost':             ['jost', 'carlito'],
    'merriweather':     ['merriweather', 'liberation serif'],
    'playfair':         ['playfair display', 'liberation serif'],
    'playfair display': ['playfair display', 'liberation serif'],
    'lora':             ['lora', 'liberation serif'],
    'pt serif':         ['pt serif', 'liberation serif'],
    'crimson':          ['crimson text', 'liberation serif'],
    'libre baskerville':['libre baskerville', 'liberation serif'],
    'baskerville':      ['libre baskerville', 'liberation serif'],
    'noto serif':       ['noto serif'],
    'bitter':           ['bitter', 'liberation serif'],
    'tinos':            ['tinos', 'liberation serif'],
    'gelasio':          ['gelasio', 'caladea'],
    'cormorant':        ['cormorant', 'eb garamond'],
    'eb garamond':      ['eb garamond'],
    'jetbrains mono':   ['jetbrains mono', 'hack', 'liberation mono'],
    'fira code':        ['fira code', 'firacode', 'hack'],
    'fira sans':        ['fira sans', 'carlito'],
    'fira mono':        ['fira mono', 'hack'],
    'source code pro':  ['source code pro', 'hack'],
    'source sans':      ['source sans pro', 'carlito'],
    'source serif':     ['source serif pro', 'liberation serif'],
    'ibm plex mono':    ['ibm plex mono', 'hack'],
    'ibm plex sans':    ['ibm plex sans', 'carlito'],
    'roboto mono':      ['roboto mono', 'hack'],
    'roboto slab':      ['roboto slab', 'liberation serif'],
    'oswald':           ['oswald', 'liberation sans'],
    'bebas neue':       ['bebas neue', 'liberation sans'],
    'lobster':          ['lobster'],
    'pacifico':         ['pacifico'],
    'dancing script':   ['dancing script'],
    'great vibes':      ['great vibes'],
    'caveat':           ['caveat'],
    'satisfy':          ['satisfy'],

    # ---- Asian scripts ----
    'ms gothic':        ['noto sans cjk jp', 'takao gothic', 'takao'],
    'ms mincho':        ['noto serif cjk jp', 'takao mincho', 'takao'],
    'meiryo':           ['noto sans cjk jp', 'takao gothic'],
    'yu gothic':        ['noto sans cjk jp'],
    'yu mincho':        ['noto serif cjk jp'],
    'ms pgothic':       ['noto sans cjk jp'],
    'ms pmincho':       ['noto serif cjk jp'],
    'simsun':           ['noto sans cjk sc', 'wqy microhei', 'wqy zenhei'],
    'simhei':           ['noto sans cjk sc', 'wqy microhei'],
    'fangsong':         ['noto serif cjk sc', 'arphic uming'],
    'kaiti':            ['noto serif cjk sc', 'arphic ukai'],
    'microsoft yahei':  ['noto sans cjk sc', 'wqy microhei'],
    'microsoft jhenghei':['noto sans cjk tc', 'wqy microhei'],
    'malgun gothic':    ['noto sans cjk kr', 'nanum gothic', 'nanum'],
    'gulim':            ['noto sans cjk kr', 'nanum'],
    'batang':           ['noto serif cjk kr', 'nanum'],
    'dotum':            ['noto sans cjk kr', 'nanum'],
    'gungsuh':          ['noto serif cjk kr', 'nanum'],

    # ---- Indic ----
    'mangal':           ['noto sans devanagari', 'lohit devanagari', 'samyak devanagari'],
    'nirmala ui':       ['noto sans devanagari', 'lohit devanagari'],
    'aparajita':        ['noto sans devanagari', 'lohit devanagari'],
    'kokila':           ['noto sans devanagari', 'lohit devanagari'],
    'utsaah':           ['noto sans devanagari', 'lohit devanagari'],
    'shruti':           ['noto sans gujarati', 'lohit gujarati', 'samyak gujarati'],
    'gautami':          ['noto sans telugu', 'lohit telugu'],
    'latha':            ['noto sans tamil', 'lohit tamil', 'samyak tamil'],
    'kartika':          ['noto sans malayalam', 'lohit malayalam', 'samyak malayalam'],
    'raavi':            ['noto sans gurmukhi', 'lohit gurmukhi'],
    'vrinda':           ['noto sans bengali', 'lohit bengali'],
    'kalinga':          ['noto sans oriya', 'lohit oriya', 'samyak oriya'],
    'sylfaen':          ['noto sans armenian'],
    'nyala':            ['noto sans ethiopic', 'abyssinica'],
    'euphemia':         ['noto sans canadian aboriginal syllabics'],

    # ---- Arabic / Hebrew ----
    'traditional arabic':['noto naskh arabic', 'amiri', 'kacst'],
    'simplified arabic':['noto naskh arabic', 'amiri', 'kacst'],
    'arabic typesetting':['noto naskh arabic', 'amiri'],
    'segoe ui arabic':  ['noto naskh arabic'],
    'aldhabi':          ['noto naskh arabic', 'amiri'],
    'urdu typesetting': ['noto naskh arabic', 'amiri'],
    'david':            ['noto sans hebrew', 'dejavu sans'],
    'miriam':           ['noto sans hebrew', 'dejavu sans'],
    'frank ruehl':      ['noto sans hebrew', 'dejavu sans'],

    # ---- Thai / Southeast Asian ----
    'leelawadee ui':    ['noto sans thai', 'loma', 'garuda'],
    'leelawadee':       ['noto sans thai', 'loma', 'garuda'],
    'tahoma thai':      ['noto sans thai', 'loma'],
    'daunpenh':         ['noto sans khmer', 'khmeros'],
    'khmer ui':         ['noto sans khmer', 'khmeros'],
    'phetsarath':       ['noto sans lao', 'lao'],
    'ishkoola':         ['noto sans sinhala', 'lklug'],
    'nirmala sinhala':  ['noto sans sinhala', 'lklug'],
    'myanmar text':     ['noto sans myanmar', 'padauk'],
}


def _is_available(font_name, installed):
    key = font_name.lower().strip()
    if not key:
        return True

    if key in installed:
        return True

    key_nospace = key.replace(' ', '')
    if key_nospace in {x.replace(' ', '') for x in installed}:
        return True

    for pattern, subs in SUBSTITUTIONS.items():
        if pattern in key:
            if not subs:
                return False
            return any(s in installed for s in subs)

    return False


def audit_docx_fonts(docx_path):
    _silence_library_noise()

    try:
        from docx import Document
    except ImportError:
        return {'fonts': [], 'available': [], 'missing': [],
                'error': 'python-docx not installed'}

    try:
        doc = Document(docx_path)
    except Exception as e:
        return {'fonts': [], 'available': [], 'missing': [], 'error': str(e)}

    fonts = set()

    try:
        for para in doc.paragraphs:
            for run in para.runs:
                if run.font.name:
                    fonts.add(run.font.name)
    except Exception:
        pass

    try:
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            if run.font.name:
                                fonts.add(run.font.name)
    except Exception:
        pass

    try:
        import re
        doc_xml = doc.element.xml
        for match in re.finditer(
            r'w:(?:ascii|hAnsi|eastAsia|cs)="([^"]+)"', doc_xml
        ):
            name = match.group(1).strip()
            if name:
                fonts.add(name)
    except Exception:
        pass

    try:
        for style in doc.styles:
            try:
                if style.font and style.font.name:
                    fonts.add(style.font.name)
            except Exception:
                continue
    except Exception:
        pass

    installed = get_installed_fonts()
    missing = sorted({f for f in fonts if not _is_available(f, installed)})
    available = sorted({f for f in fonts if f not in missing})

    log(f"Found {len(fonts)} font(s); {len(missing)} will be substituted")

    return {
        'fonts': sorted(fonts),
        'available': available,
        'missing': missing,
        'installed_count': len(installed),
    }


if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.stderr.write("Usage: audit_docx_fonts.py <input.docx>\n")
        _emit({'fonts': [], 'available': [], 'missing': [],
               'error': 'no input path provided'})
        sys.exit(1)
    try:
        result = audit_docx_fonts(sys.argv[1])
        _emit(result)
        sys.exit(0)
    except Exception as e:
        log(f"FATAL: {e}")
        _emit({'fonts': [], 'available': [], 'missing': [], 'error': str(e)})
        sys.exit(0)