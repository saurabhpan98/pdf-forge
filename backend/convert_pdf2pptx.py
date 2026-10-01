"""
High-fidelity PDF → PowerPoint conversion.

Two output modes:

  • 'image' — Each PDF page is rendered at a target DPI and placed as a
    full-slide image. Preserves visual fidelity exactly (what you see in
    the PDF is what appears on the slide). Text is NOT editable.

  • 'text'  — Text blocks and images are extracted as separate editable
    objects. Text is fully editable in PowerPoint. Positions and font
    sizes are preserved from the PDF; fonts are mapped to installed
    equivalents. Visual fidelity is 70–85% depending on the source.

Both modes use PyMuPDF for extraction and python-pptx for output.

Emits a marker-prefixed JSON result on stdout:
    __RESULT__{"success":true,"mode":"text","slides":5,...}
"""
import sys
import os
import io
import json

try:
    import fitz  # PyMuPDF
except ImportError:
    try:
        import pymupdf as fitz
    except ImportError as exc:
        sys.stderr.write("[pptx] FATAL: PyMuPDF not installed.\n")
        raise exc

try:
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
except ImportError as exc:
    sys.stderr.write("[pptx] FATAL: python-pptx not installed.\n")
    raise exc

MARKER = '__RESULT__'

# PDF points to EMU: 1 pt = 1/72 inch, 1 inch = 914400 EMU
PT_TO_EMU = 12700


def log(msg):
    sys.stderr.write(f"[pptx] {msg}\n")


def _emit(payload):
    sys.stdout.write(f"{MARKER}{json.dumps(payload)}\n")
    sys.stdout.flush()


# ===========================================================================
# Font mapping
# ===========================================================================
FONT_ALIASES = {
    'helvetica': 'Arial',
    'liberationsans': 'Arial',
    'liberationserif': 'Times New Roman',
    'liberationmono': 'Courier New',
    'arimo': 'Arial',
    'tinos': 'Times New Roman',
    'cousine': 'Courier New',
    'carlito': 'Calibri',
    'caladea': 'Cambria',
    'nimbussans': 'Arial',
    'dejavusans': 'Arial',
    'dejavuserif': 'Times New Roman',
    'dejavusansmono': 'Courier New',
    'notosans': 'Calibri',
    'notoserif': 'Cambria',
    'opensans': 'Segoe UI',
    'lato': 'Segoe UI',
    'roboto': 'Segoe UI',
    'inter': 'Segoe UI',
}


def _map_font(raw_name):
    if not raw_name:
        return 'Calibri'
    name = raw_name
    if '+' in name:
        name = name.split('+', 1)[1]
    for suffix in (
        '-BoldItalic', '-BoldOblique', '-Bold', '-Italic', '-Oblique',
        ' Bold Italic', ' Bold', ' Italic', ' Oblique',
        'BoldItalic', 'Bold', 'Italic', 'Oblique',
        '-Regular', ' Regular', '-Roman', ' Roman',
        'MT', 'PS', 'MS', 'Std', 'Pro',
    ):
        if name.endswith(suffix):
            name = name[: -len(suffix)]
            break
    name = name.strip(' -_')
    if not name:
        return 'Calibri'
    low = name.lower().replace(' ', '').replace('-', '').replace('_', '')
    if low in FONT_ALIASES:
        return FONT_ALIASES[low]
    return name


def _extract_color(span):
    c = span.get('color', 0)
    if isinstance(c, int):
        return ((c >> 16) & 0xFF, (c >> 8) & 0xFF, c & 0xFF)
    if isinstance(c, (list, tuple)) and len(c) >= 3:
        return tuple(int(v) for v in c[:3])
    return (0, 0, 0)


# ===========================================================================
# IMAGE MODE
# ===========================================================================
def _raster_mode(input_path, output_path, options):
    dpi = max(72, min(300, int(options.get('dpi', 150))))
    quality = max(50, min(95, int(options.get('quality', 85))))

    src = fitz.open(input_path)
    prs = Presentation()

    # Slide dimensions match the first page
    first = src[0]
    rect = first.rect
    prs.slide_width = Emu(int(rect.width * PT_TO_EMU))
    prs.slide_height = Emu(int(rect.height * PT_TO_EMU))

    blank_layout = prs.slide_layouts[6]

    scale = dpi / 72.0
    mat = fitz.Matrix(scale, scale)

    slide_count = 0
    for page in src:
        pix = page.get_pixmap(matrix=mat, alpha=False, colorspace=fitz.csRGB)
        try:
            img_bytes = pix.tobytes(output='jpeg', jpg_quality=quality)
        except TypeError:
            img_bytes = pix.tobytes(output='jpeg')

        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            io.BytesIO(img_bytes),
            0, 0,
            width=prs.slide_width,
            height=prs.slide_height,
        )
        slide_count += 1

    src.close()
    prs.save(output_path)
    return {
        'mode': 'image',
        'slides': slide_count,
        'dpi': dpi,
        'quality': quality,
    }


# ===========================================================================
# TEXT MODE
# ===========================================================================
def _text_mode(input_path, output_path, options):
    min_font = float(options.get('minFontSize', 4.0))
    min_block_area = float(options.get('minBlockArea', 20.0))  # square points

    src = fitz.open(input_path)
    prs = Presentation()

    first = src[0]
    rect = first.rect
    prs.slide_width = Emu(int(rect.width * PT_TO_EMU))
    prs.slide_height = Emu(int(rect.height * PT_TO_EMU))

    blank_layout = prs.slide_layouts[6]

    slides_created = 0
    text_boxes = 0
    images = 0

    for page in src:
        slide = prs.slides.add_slide(blank_layout)
        slides_created += 1

        # ---- Images first (so text sits on top) ----
        try:
            img_info = page.get_image_info(xrefs=True)
        except Exception:
            img_info = []

        for info in img_info:
            bbox = info.get('bbox')
            if not bbox:
                continue
            x0, y0, x1, y1 = bbox
            w = x1 - x0
            h = y1 - y0
            if w < 4 or h < 4:
                continue

            xref = info.get('xref', 0)
            if not xref:
                continue

            try:
                img_dict = src.extract_image(xref)
                if not img_dict:
                    continue
                img_bytes = img_dict.get('image')
                if not img_bytes:
                    continue

                slide.shapes.add_picture(
                    io.BytesIO(img_bytes),
                    Emu(int(x0 * PT_TO_EMU)),
                    Emu(int(y0 * PT_TO_EMU)),
                    width=Emu(int(w * PT_TO_EMU)),
                    height=Emu(int(h * PT_TO_EMU)),
                )
                images += 1
            except Exception as ex:
                log(f"image insert failed: {ex}")

        # ---- Text blocks ----
        try:
            text_dict = page.get_text('dict')
        except Exception:
            text_dict = {}

        blocks = text_dict.get('blocks', [])

        for block in blocks:
            if block.get('type') != 0:
                continue

            bbox = block.get('bbox')
            if not bbox:
                continue
            x0, y0, x1, y1 = bbox
            w = x1 - x0
            h = y1 - y0
            if w < 4 or h < 4:
                continue
            if w * h < min_block_area:
                continue

            # Collect all lines in this block
            lines = []
            for line in block.get('lines', []):
                spans = line.get('spans', [])
                if not spans:
                    continue

                line_parts = []
                for span in spans:
                    t = span.get('text', '')
                    if not t:
                        continue
                    size = float(span.get('size', 12))
                    if size < min_font:
                        continue

                    flags = span.get('flags', 0)
                    line_parts.append({
                        'text': t,
                        'size': size,
                        'font': span.get('font', ''),
                        'bold': bool(flags & 16),
                        'italic': bool(flags & 2),
                        'color': _extract_color(span),
                    })

                if not line_parts:
                    continue

                lines.append(line_parts)

            if not lines:
                continue

            # Create a textbox
            tb = slide.shapes.add_textbox(
                Emu(int(x0 * PT_TO_EMU)),
                Emu(int(y0 * PT_TO_EMU)),
                Emu(int(w * PT_TO_EMU)),
                Emu(int(h * PT_TO_EMU)),
            )
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = 0
            tf.margin_right = 0
            tf.margin_top = 0
            tf.margin_bottom = 0

            for l_idx, line_spans in enumerate(lines):
                if l_idx == 0:
                    p = tf.paragraphs[0]
                else:
                    p = tf.add_paragraph()
                p.space_before = Pt(0)
                p.space_after = Pt(0)

                for span in line_spans:
                    run = p.add_run()
                    run.text = span['text']
                    if span['font']:
                        run.font.name = _map_font(span['font'])
                    run.font.size = Pt(span['size'])
                    run.font.bold = span['bold']
                    run.font.italic = span['italic']
                    try:
                        run.font.color.rgb = RGBColor(*span['color'])
                    except Exception:
                        pass

            text_boxes += 1

    src.close()
    prs.save(output_path)

    return {
        'mode': 'text',
        'slides': slides_created,
        'textBoxes': text_boxes,
        'images': images,
    }


# ===========================================================================
# Entry
# ===========================================================================
def convert(input_path, output_path, options=None):
    options = options or {}
    mode = options.get('mode', 'image')

    if mode not in ('image', 'text'):
        mode = 'image'

    log(f"Mode: {mode}, input size: {os.path.getsize(input_path)} bytes")

    if mode == 'text':
        result = _text_mode(input_path, output_path, options)
    else:
        result = _raster_mode(input_path, output_path, options)

    result['success'] = True
    result['sizeBytes'] = os.path.getsize(output_path)
    return result


if __name__ == '__main__':
    if len(sys.argv) < 3:
        sys.stderr.write("Usage: convert_pdf2pptx.py <input.pdf> <output.pptx> [payload.json]\n")
        _emit({'success': False, 'error': 'missing arguments'})
        sys.exit(1)

    in_pdf = sys.argv[1]
    out_pptx = sys.argv[2]

    options = {}
    if len(sys.argv) > 3 and os.path.exists(sys.argv[3]):
        try:
            with open(sys.argv[3], 'r', encoding='utf-8') as f:
                options = json.load(f)
        except Exception as e:
            log(f"payload parse failed: {e}")

    try:
        result = convert(in_pdf, out_pptx, options)
        _emit(result)
        sys.exit(0)
    except Exception as e:
        log(f"FATAL: {e}")
        import traceback
        traceback.print_exc(file=sys.stderr)
        _emit({'success': False, 'error': str(e)})
        sys.exit(1)